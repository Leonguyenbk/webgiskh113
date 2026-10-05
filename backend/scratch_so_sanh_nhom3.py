"""So sánh danh sách thửa cần thu thập Nhóm 3 (DANHSACHNHOM3BMT.xlsx) với
dữ liệu hiện có: thửa đã thuộc Nhóm 1/2 (dong_bo_du_lieu.phan_loai_ke_hoach_2959)
và thửa đã thu thập (du_lieu_gcn). Kết quả in ra số liệu tổng hợp + xuất file
Excel danh sách thửa CÒN THIẾU dữ liệu (chưa thuộc Nhóm 1/2, chưa thu thập).

Chạy 1 lần, không phải route/service thường trực — xoá sau khi dùng xong.
"""
from __future__ import annotations

import openpyxl
import psycopg2
import psycopg2.extras

from app.config import Config

SRC = r"D:\QUANGTUAN\webgiskh113\DANHSACHNHOM3BMT.xlsx"
OUT = r"D:\QUANGTUAN\webgiskh113\NHOM3_CON_THIEU_DU_LIEU.xlsx"

NHOM_1_2 = {"NHÓM 1", "NHÓM 2"}


def doc_danh_sach():
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    ws = wb["Sheet1"]
    rows = ws.iter_rows(values_only=True)
    next(rows)  # header
    out = []
    for ma_xa, so_to, so_thua, ma_dinh_danh in rows:
        if ma_xa is None or so_to is None or so_thua is None:
            continue
        out.append((str(int(ma_xa)), int(so_to), int(so_thua), ma_dinh_danh))
    wb.close()
    return out


def main():
    danh_sach = doc_danh_sach()
    print(f"Đọc {len(danh_sach)} dòng từ {SRC}")

    conn = psycopg2.connect(Config.DATABASE_URL)
    conn.autocommit = False
    cur = conn.cursor()

    cur.execute(
        """
        create temporary table tmp_nhom3 (
            ma_xa text, so_to integer, so_thua integer, ma_dinh_danh text
        ) on commit drop
        """
    )
    psycopg2.extras.execute_values(
        cur,
        "insert into tmp_nhom3 (ma_xa, so_to, so_thua, ma_dinh_danh) values %s",
        danh_sach,
        page_size=5000,
    )
    cur.execute("create index on tmp_nhom3 (ma_xa, so_to, so_thua)")

    # Trùng trong chính file nguồn?
    cur.execute(
        """
        select count(*) from (
            select ma_xa, so_to, so_thua from tmp_nhom3
            group by ma_xa, so_to, so_thua having count(*) > 1
        ) d
        """
    )
    (so_dong_trung,) = cur.fetchone()

    cur.execute("select count(distinct (ma_xa, so_to, so_thua)) from tmp_nhom3")
    (tong_thua_duy_nhat,) = cur.fetchone()

    # Có thật trong thua_dat không (đã nhập GML)?
    cur.execute(
        """
        select count(distinct (t.ma_xa, t.so_to, t.so_thua))
        from tmp_nhom3 t
        join thua_dat p on p.ma_xa = t.ma_xa and p.so_to = t.so_to and p.so_thua = t.so_thua
        """
    )
    (co_that_thua_dat,) = cur.fetchone()

    # Đã thuộc Nhóm 1/2 (KH 2959)?
    cur.execute(
        """
        select count(distinct (t.ma_xa, t.so_to, t.so_thua))
        from tmp_nhom3 t
        join dong_bo_du_lieu d
          on d.ma_xa = t.ma_xa and d.so_to = t.so_to and d.so_thua = t.so_thua
        where upper(trim(d.phan_loai_ke_hoach_2959)) in ('NHÓM 1', 'NHÓM 2')
        """
    )
    (da_nhom_1_2,) = cur.fetchone()

    # Đã thu thập (có trong du_lieu_gcn, mọi nguồn)?
    cur.execute(
        """
        select count(distinct (t.ma_xa, t.so_to, t.so_thua))
        from tmp_nhom3 t
        join du_lieu_gcn g
          on g.madvhc_soto_sothua = t.ma_xa || '_' || t.so_to::text || '_' || t.so_thua::text
        """
    )
    (da_thu_thap,) = cur.fetchone()

    # CÒN THIẾU: chưa Nhóm 1/2 VÀ chưa thu thập. Kèm cờ "có bản đồ" (thua_dat)
    # để tách thửa thật sự chờ nhập dữ liệu khỏi thửa chưa số hóa GML.
    cur.execute(
        """
        select distinct t.ma_xa, t.so_to, t.so_thua, t.ma_dinh_danh,
               (p.id is not null) as co_ban_do
        from tmp_nhom3 t
        left join thua_dat p on p.ma_xa = t.ma_xa and p.so_to = t.so_to and p.so_thua = t.so_thua
        where not exists (
            select 1 from dong_bo_du_lieu d
            where d.ma_xa = t.ma_xa and d.so_to = t.so_to and d.so_thua = t.so_thua
              and upper(trim(d.phan_loai_ke_hoach_2959)) in ('NHÓM 1', 'NHÓM 2')
        )
        and not exists (
            select 1 from du_lieu_gcn g
            where g.madvhc_soto_sothua = t.ma_xa || '_' || t.so_to::text || '_' || t.so_thua::text
        )
        order by t.ma_xa, t.so_to, t.so_thua
        """
    )
    con_thieu = cur.fetchall()
    co_ban_do = sum(1 for r in con_thieu if r[4])
    khong_co_ban_do = len(con_thieu) - co_ban_do

    conn.rollback()  # temp table on commit drop, nhưng không cần giữ transaction
    cur.close()
    conn.close()

    print(f"Số thửa duy nhất trong file nguồn: {tong_thua_duy_nhat} (dòng trùng: {so_dong_trung})")
    print(f"Có thật trong thua_dat (đã nhập GML): {co_that_thua_dat}")
    print(f"Đã thuộc Nhóm 1/2 (KH 2959): {da_nhom_1_2}")
    print(f"Đã thu thập (có trong du_lieu_gcn): {da_thu_thap}")
    print(f"CÒN THIẾU dữ liệu (chưa Nhóm 1/2, chưa thu thập): {len(con_thieu)}")
    print(f"  - Có trên bản đồ (thua_dat), chỉ thiếu dữ liệu thuộc tính: {co_ban_do}")
    print(f"  - KHÔNG có trên bản đồ (chưa nhập GML) — cần kiểm tra riêng: {khong_co_ban_do}")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Con thieu du lieu"
    ws.append(["MAXA", "SOTO", "SOTHUA", "MADINHDANH", "CO_BAN_DO"])
    for ma_xa, so_to, so_thua, ma_dinh_danh, co_bd in con_thieu:
        ws.append([ma_xa, so_to, so_thua, ma_dinh_danh, "Có" if co_bd else "Không"])
    wb.save(OUT)
    print(f"Đã ghi {len(con_thieu)} dòng ra {OUT}")


if __name__ == "__main__":
    main()
